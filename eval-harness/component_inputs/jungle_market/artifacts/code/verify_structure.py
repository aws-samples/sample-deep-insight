
from docx import Document

for path in ['./artifacts/final_report_with_citations.docx', './artifacts/final_report.docx']:
    doc = Document(path)
    print(f"=== {path} ===")
    print("Paragraphs:", len(doc.paragraphs))
    print("Tables:", len(doc.tables))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    for h in headings:
        print(" -", h)
    print()
